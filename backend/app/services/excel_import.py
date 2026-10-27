"""Excel import engine shared by PO, GRN, wholesale bill, item master and price update imports.

preview(): read the .xlsx (codes kept as text), validate every row against the database, store an import job (PREVIEW).
commit():  re-validate (stock / prices may have changed), then create the documents through the EXISTING services —
           PO (create_po, approval flow unchanged), GRN (create_purchase, draft or posted), wholesale bill (create_invoice:
           server pricing, GST, discount policy, stock, numbering), items (create_product / update, never stock),
           prices (PRICE_CHANGE approval when switched on, history + store-wise propagation).
A job commits once (row lock); the same file (hash + header) cannot create a second PO / GRN / bill unless the user says so.
Invalid rows are never skipped silently: commit refuses them unless the user explicitly imports only the valid rows.
"""
import hashlib
import io
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Optional

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import text

from app.services.approval_service import handler

MAX_ROWS, MAX_BYTES = 5000, 5 * 1024 * 1024
TEXT_KEYS = {"item_code", "barcode", "batch_no", "hsn", "supplier_code", "customer_code", "po_no"}
PRICE_FIELDS = {"basic_cost": "basic_cost", "cp": "cost_price", "mrp": "mrp", "sp": "selling_price", "wsp": "wsp"}

# column key, header, required, kind (text | num | date | bool)
ITEM = [("item_code", "Item Code", False, "text"), ("barcode", "Barcode/EAN", False, "text"), ("item_name", "Item Name", False, "text")]
COLUMNS = {
    "PO": ITEM + [("unit", "Unit", False, "text"), ("qty", "Quantity", True, "num"), ("basic_cost", "Basic Cost", True, "num"),
                  ("mrp", "MRP", False, "num"), ("discount_pct", "Discount %", False, "num"), ("gst_pct", "GST %", False, "num"),
                  ("remarks", "Remarks", False, "text")],
    "GRN": ITEM + [("unit", "Unit", False, "text"), ("received_qty", "Received Qty", True, "num"), ("accepted_qty", "Accepted Qty", False, "num"),
                   ("rejected_qty", "Rejected Qty", False, "num"), ("free_qty", "Free Qty", False, "num"), ("basic_cost", "Basic Cost", True, "num"),
                   ("mrp", "MRP", False, "num"), ("sp", "SP", False, "num"), ("discount_pct", "Discount %", False, "num"),
                   ("gst_pct", "GST %", False, "num"), ("batch_no", "Batch Number", False, "text"), ("mfg_date", "Mfg Date", False, "date"),
                   ("expiry_date", "Expiry Date", False, "date"), ("remarks", "Remarks", False, "text")],
    "WHOLESALE": ITEM + [("unit", "Unit", False, "text"), ("qty", "Quantity", True, "num"), ("rate", "Rate (blank = WSP / SP)", False, "num"),
                         ("discount_pct", "Discount %", False, "num"), ("batch_no", "Batch Number", False, "text"), ("remarks", "Line Remarks", False, "text")],
    "ITEMS": [("item_code", "Item Code", False, "text"), ("item_name", "Item Name", True, "text"), ("barcode", "Barcode/EAN", False, "text"),
              ("item_type", "Item Type (FG/RM/PM/SFG/CONS/OTH)", False, "text"), ("category", "Category", True, "text"), ("brand", "Brand", False, "text"), ("hsn", "HSN", False, "text"), ("gst_pct", "GST %", False, "num"),
              ("unit", "Unit", False, "text"), ("basic_cost", "Basic Cost", False, "num"), ("cp", "CP", False, "num"), ("mrp", "MRP", False, "num"),
              ("sp", "SP", False, "num"), ("wsp", "WSP", False, "num"), ("batch_tracking", "Batch Tracking (Y/N)", False, "bool"),
              ("active", "Active (Y/N)", False, "bool")],
    "PRICES": [("item_code", "Item Code", False, "text"), ("barcode", "Barcode/EAN", False, "text"), ("item_name", "Item Name", False, "text"),
               ("basic_cost", "Basic Cost", False, "num"), ("cp", "CP", False, "num"), ("mrp", "MRP", False, "num"), ("sp", "SP", False, "num"),
               ("wsp", "WSP", False, "num")],
}
EXAMPLE = {"PO": ["ITM001", "8901234567890", "Example item", "PCS", 10, 45.5, 60, 0, 5, ""],
           "GRN": ["ITM001", "", "Example item", "PCS", 10, 10, 0, 0, 45.5, 60, 58, 0, 5, "B2410", "2026-09-01", "2027-08-31", ""],
           "WHOLESALE": ["ITM001", "", "Example item", "PCS", 24, "", 0, "", ""],
           "ITEMS": ["", "Example item 500g", "8901234567890", "FG", "Grocery", "Brand X", "0902", 5, "PCS", 40, 42, 60, 55, 50, "N", "Y"],
           "PRICES": ["ITM001", "", "Example item", 40, 42, 60, 55, 50]}


def template(module: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = module.lower()
    cols = COLUMNS[module]
    for i, (key, head, req, kind) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=head + (" *" if req else ""))
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4F46E5" if req else "64748B")
        ws.column_dimensions[c.column_letter].width = max(12, len(head) + 4)
        ex = ws.cell(row=2, column=i, value=EXAMPLE[module][i - 1] if i - 1 < len(EXAMPLE[module]) else None)
        if key in TEXT_KEYS:
            ex.number_format = "@"
            for r in range(3, 2003):
                ws.cell(row=r, column=i).number_format = "@"     # keep leading zeros in codes / barcodes
    info = wb.create_sheet("Instructions")
    for r, line in enumerate([
        f"{module} import — fill row 2 onwards (delete the example row). Columns marked * are required.",
        "Identify each item by Item Code or Barcode/EAN (kept as text — leading zeros are safe).",
        "Dates: YYYY-MM-DD or DD-MM-YYYY. Numbers: plain numbers, no currency signs.",
        "Upload on the Excel Import screen, check the preview, correct any red rows and upload again, then confirm.",
        "Totals, GST and stock are always calculated by the server — Excel amounts are not trusted.",
        f"At most {MAX_ROWS} rows per file.",
    ], 1):
        info.cell(row=r, column=1, value=line)
    info.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _norm(h) -> str:
    return re.sub(r"[^a-z0-9%]+", " ", str(h or "").lower().replace("*", "")).strip()


def _cell(v, kind: str):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if kind == "text":
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        return str(v).strip()
    if kind == "num":
        try:
            return Decimal(str(v).replace(",", "").strip())
        except InvalidOperation:
            raise ValueError(f"'{v}' is not a number")
    if kind == "date":
        if isinstance(v, datetime):
            return v.date()
        if isinstance(v, date):
            return v
        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y"):
            try:
                return datetime.strptime(str(v).strip(), fmt).date()
            except ValueError:
                pass
        raise ValueError(f"'{v}' is not a date (use YYYY-MM-DD)")
    if kind == "bool":
        return str(v).strip().lower() in ("y", "yes", "true", "1")
    return v


def parse(module: str, data: bytes) -> list[dict]:
    if len(data) > MAX_BYTES:
        raise HTTPException(400, "File is larger than 5 MB.")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "Not a valid .xlsx file — download the template and save it as Excel Workbook (.xlsx).")
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = next(it, None) or ()
    cols = COLUMNS[module]
    want = {_norm(h): (k, kind) for k, h, _, kind in cols}
    idx = {}
    for i, h in enumerate(head):
        n = _norm(h)
        if n in want:
            idx[i] = want[n]
    missing = [h for k, h, req, _ in cols if req and k not in {v[0] for v in idx.values()}]
    if missing:
        raise HTTPException(400, f"Missing column(s): {', '.join(missing)}. Use the template headings.")
    rows = []
    for n, raw in enumerate(it, start=2):
        if raw is None or all(v is None or str(v).strip() == "" for v in raw):
            continue
        rec, errs = {"_row": n}, []
        for i, (k, kind) in idx.items():
            try:
                rec[k] = _cell(raw[i] if i < len(raw) else None, kind)
            except ValueError as e:
                errs.append(f"{k.replace('_', ' ')}: {e}")
        rec["_errors"] = errs
        rows.append(rec)
        if len(rows) > MAX_ROWS:
            raise HTTPException(400, f"More than {MAX_ROWS} rows — split the file.")
    if not rows:
        raise HTTPException(400, "The file has no data rows.")
    return rows


# ── validation ────────────────────────────────────────────────────────────────
async def _items(db, rows: list[dict]) -> dict:
    """Resolve rows to products by item code, else barcode. Returns {row_no: product dict}."""
    codes = list({r["item_code"].upper() for r in rows if r.get("item_code")})
    bars = list({r["barcode"] for r in rows if r.get("barcode")})
    by_code = {r["k"]: dict(r) for r in (await db.execute(text("""SELECT UPPER(item_code) AS k, * FROM products WHERE UPPER(item_code) = ANY(:c)"""),
                                                          {"c": codes or [""]})).mappings()}
    by_bar = {r["k"]: dict(r) for r in (await db.execute(text("""SELECT pb.barcode AS k, p.* FROM product_barcodes pb JOIN products p ON p.id = pb.product_id
        WHERE pb.barcode = ANY(:b) UNION SELECT p.barcode AS k, p.* FROM products p WHERE p.barcode = ANY(:b)"""), {"b": bars or [""]})).mappings()}
    out = {}
    for r in rows:
        p = by_code.get((r.get("item_code") or "").upper()) or by_bar.get(r.get("barcode") or "")
        if p:
            out[r["_row"]] = p
    return out


def _d(v) -> Decimal:
    return Decimal(str(v or 0))


async def validate(db, module: str, rows: list[dict], header: dict, user) -> list[dict]:
    from app.core import access
    if module in ("PO", "GRN", "WHOLESALE"):
        if not header.get("outlet_id"):
            raise HTTPException(400, "Choose the store.")
        access.check_outlet(int(header["outlet_id"]))
    prods = await _items(db, rows) if module != "ITEMS" else {}
    outside: set = set()
    if prods:                                    # same data-scope rule the services enforce, shown per row
        try:
            await access.check_products(db, list({p["id"] for p in prods.values()}))
        except HTTPException:
            for pid in {p["id"] for p in prods.values()}:
                try:
                    await access.check_products(db, [pid])
                except HTTPException:
                    outside.add(pid)
    seen: dict = {}
    po_lines = {}
    if module == "GRN" and header.get("po_id"):
        po_lines = {r["product_id"]: dict(r) for r in (await db.execute(text("""SELECT id, product_id, COALESCE(NULLIF(order_qty, 0), qty) AS ordered,
            received_qty FROM po_items WHERE po_id = :p"""), {"p": header["po_id"]})).mappings()}
    today = date.today()
    for r in rows:
        e = r["_errors"] = list(r.get("_errors") or [])
        p = prods.get(r["_row"])
        if module != "ITEMS":
            if not (r.get("item_code") or r.get("barcode")):
                e.append("Item Code or Barcode is required")
            elif not p:
                e.append("Unknown item code / barcode")
            else:
                r["product_id"], r["item_name"], r["_unit"] = p["id"], p["name"], (p.get("unit") or "PCS").upper()
                if p["id"] in outside:
                    e.append("Item is outside your categories / groups")
                if r.get("unit") and r["unit"].upper() not in (r["_unit"], (p.get("purchase_unit") or "").upper()):
                    e.append(f"Unit {r['unit']} is not this item's unit ({r['_unit']})")
        for k, _, req, kind in COLUMNS[module]:
            if req and r.get(k) in (None, ""):
                e.append(f"{k.replace('_', ' ')} is missing")
            if kind == "num" and r.get(k) is not None and r[k] < 0:
                e.append(f"{k.replace('_', ' ')} cannot be negative")
        if module == "ITEMS":                       # same item twice in one file
            key = (r.get("item_code") or r.get("item_name") or "").upper() or None
        elif module in ("GRN", "WHOLESALE"):        # the same item may come in several batches
            key = (r["product_id"], r.get("batch_no") or "") if r.get("product_id") else None
        else:
            key = r.get("product_id")
        if key:
            if key in seen:
                e.append(f"Duplicate of row {seen[key]}")
            seen.setdefault(key, r["_row"])
        if module == "PO" and r.get("qty") is not None and r["qty"] <= 0:
            e.append("Quantity must be more than zero")
        if module == "GRN" and p:
            rec, acc, rej = _d(r.get("received_qty")), r.get("accepted_qty"), _d(r.get("rejected_qty"))
            acc = rec - rej if acc is None else acc
            r["accepted_qty"] = acc
            if acc + rej != rec:
                e.append(f"Accepted {acc} + rejected {rej} must equal received {rec}")
            if acc <= 0 and rec > 0 and rej == 0:
                e.append("Accepted quantity is zero")
            needs_batch = p.get("has_batch_tracking") or p.get("is_batch_required")
            if needs_batch and not r.get("batch_no"):
                e.append("This item needs a batch number")
            if needs_batch and not r.get("expiry_date"):
                e.append("This item needs an expiry date")
            if r.get("expiry_date") and r["expiry_date"] <= today:
                e.append(f"Batch already expired ({r['expiry_date']})")
            if r.get("mfg_date") and r.get("expiry_date") and r["mfg_date"] >= r["expiry_date"]:
                e.append("Expiry must be after the manufacturing date")
            if po_lines:
                pl = po_lines.get(p["id"])
                if not pl:
                    e.append("Item is not on the selected PO")
                else:
                    r["po_item_id"] = pl["id"]
                    left = _d(pl["ordered"]) - _d(pl["received_qty"])
                    if acc > left:
                        if header.get("allow_over_receipt") and header.get("_can_approve"):
                            r["_info"] = f"Over-receipt by {acc - left} (allowed by approver)"
                        else:
                            e.append(f"Over-receipt: PO balance is {left}, accepted {acc} (needs a PO approver to allow)")
        if module == "WHOLESALE" and p:
            if r.get("qty") is not None and r["qty"] <= 0:
                e.append("Quantity must be more than zero")
            wsp, sp = _d(p.get("wsp")), _d(p.get("selling_price"))
            default = wsp if wsp > 0 else sp
            if r.get("rate") is None:
                r["rate"], r["_price_source"] = default, "WSP" if wsp > 0 else "SP"
                if default <= 0:
                    e.append("Item has no WSP / SP — enter a rate")
            elif r["rate"] != default:
                if not header.get("_can_change_rate"):
                    e.append(f"Rate {r['rate']} differs from the price list ({default}); only users who may change prices can override it")
                r["_price_source"] = "Manual"
            if p.get("mrp") and r.get("rate") and _d(r["rate"]) > _d(p["mrp"]) and _d(p["mrp"]) > 0:
                e.append(f"Rate above MRP ({p['mrp']})")
            if header.get("outlet_id"):
                have = _d((await db.execute(text("SELECT stock_qty FROM outlet_stock WHERE outlet_id = :o AND product_id = :p"),
                                            {"o": int(header["outlet_id"]), "p": p["id"]})).scalar())
                if have < _d(r.get("qty")) and not p.get("allow_neg_stock"):
                    e.append(f"Only {have} in stock at this store")
        if module == "ITEMS":
            await _validate_item(db, r, header)
        if module == "PRICES" and p:
            changes = {}
            for col, field in PRICE_FIELDS.items():
                if r.get(col) is not None and _d(r[col]) != _d(p.get(field)):
                    changes[field] = {"old": str(_d(p.get(field))), "new": str(r[col])}
            mrp = r.get("mrp") if r.get("mrp") is not None else _d(p.get("mrp"))
            for f in ("sp", "wsp"):
                if r.get(f) is not None and mrp and _d(r[f]) > _d(mrp) and _d(mrp) > 0:
                    e.append(f"{f.upper()} {r[f]} is above MRP {mrp}")
            if r.get("cp") is not None and r.get("sp") is not None and r["sp"] < r["cp"]:
                r["_info"] = "Warning: SP below CP"
            r["changes"] = changes
            r["_status_hint"] = "UNCHANGED" if not changes else "CHANGED"
        r["status"] = "INVALID" if e else (r.get("_status_hint") or "VALID")
    return rows


async def _validate_item(db, r: dict, header: dict):
    e = r["_errors"]
    cat = (await db.execute(text("SELECT id FROM item_categories WHERE LOWER(name) = LOWER(:n) LIMIT 1"), {"n": r.get("category") or ""})).scalar()
    if r.get("category") and not cat:
        e.append(f"Unknown category '{r['category']}' (create it in Masters first)")
    r["category_id"] = cat
    if r.get("brand"):
        b = (await db.execute(text("SELECT id FROM brands WHERE LOWER(name) = LOWER(:n) LIMIT 1"), {"n": r["brand"]})).scalar()
        if not b:
            e.append(f"Unknown brand '{r['brand']}'")
        r["brand_id"] = b
    if r.get("item_type") and r["item_type"].upper() not in ("FG", "RM", "PM", "SFG", "CONS", "OTH"):
        e.append("Item Type must be FG, RM, PM, SFG, CONS or OTH")
    if r.get("gst_pct") is not None and r["gst_pct"] not in {Decimal(x) for x in ("0", "0.25", "3", "5", "12", "18", "28", "40")}:
        e.append(f"GST {r['gst_pct']}% is not a valid rate")
    ex = None
    if r.get("item_code"):
        ex = (await db.execute(text("SELECT id, name, mrp, selling_price, wsp, cost_price, basic_cost FROM products WHERE UPPER(item_code) = UPPER(:c)"),
                               {"c": r["item_code"]})).mappings().first()
    if ex:
        r["product_id"], r["mode"] = ex["id"], "UPDATE"
        if not header.get("_can_edit"):
            e.append("Item exists — you may only create new items (no Item Master edit permission)")
        if any(r.get(c) is not None and _d(r[c]) != _d(ex[f]) for c, f in PRICE_FIELDS.items()) and not header.get("_can_change_price"):
            e.append("Price change on an existing item needs the 'change price' permission (use the Price Update import)")
    else:
        r["mode"] = "CREATE"
        if r.get("item_name") and (await db.execute(text("SELECT 1 FROM products WHERE LOWER(TRIM(name)) = LOWER(TRIM(:n))"), {"n": r["item_name"]})).first():
            e.append("An item with this name already exists")
    if r.get("barcode"):
        hit = (await db.execute(text("""SELECT p.id FROM products p WHERE p.barcode = :b UNION SELECT pb.product_id FROM product_barcodes pb WHERE pb.barcode = :b LIMIT 1"""),
                                {"b": r["barcode"]})).scalar()
        if hit and hit != r.get("product_id"):
            e.append(f"Barcode {r['barcode']} is already on another item")
    if r.get("mrp") and r.get("sp") and r["sp"] > r["mrp"]:
        e.append("SP is above MRP")


def summary(rows: list[dict]) -> dict:
    s = {"total": len(rows), "valid": 0, "invalid": 0, "changed": 0, "unchanged": 0}
    for r in rows:
        s[r["status"].lower()] = s.get(r["status"].lower(), 0) + 1
    s["valid"] = s["total"] - s["invalid"]
    return s


def jsonable(rows):
    return json.loads(json.dumps(rows, default=lambda v: str(v) if isinstance(v, (Decimal, date)) else None))


def file_hash(data: bytes, header: dict) -> str:
    return hashlib.sha256(data + json.dumps({k: v for k, v in header.items() if not k.startswith("_")}, sort_keys=True, default=str).encode()).hexdigest()


def errors_xlsx(module: str, rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    cols = [k for k, *_ in COLUMNS[module]]
    ws.append(["Excel Row", *[h for _, h, *_ in COLUMNS[module]], "Status", "Error"])
    for r in rows:
        if r["status"] == "INVALID" or r.get("_failed"):
            ws.append([r["_row"], *[r.get(k) for k in cols], r["status"], "; ".join(r.get("_errors") or []) or r.get("_failed")])
    for c in ws[1]:
        c.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── price change approval (bulk) ──────────────────────────────────────────────
async def apply_prices(db, changes: list[dict], user_id: int, source: str, ref: str, reason: Optional[str]) -> int:
    from app.services.product_service import propagate_product_prices
    n = 0
    for c in changes:
        sets, p = [], {"i": int(c["product_id"])}
        for field, v in c["changes"].items():
            cur = (await db.execute(text(f"SELECT {field} FROM products WHERE id = :i FOR UPDATE"), p)).scalar()
            sets.append(f"{field} = :{field}")
            p[field] = Decimal(v["new"])
            await db.execute(text("""INSERT INTO price_change_log (product_id, field, old_value, new_value, source, ref, reason, changed_by)
                VALUES (:i, :f, :o, :n, :s, :r, :why, :u)"""), {"i": p["i"], "f": field, "o": cur, "n": p[field], "s": source, "r": ref, "why": reason, "u": user_id})
        if sets:
            await db.execute(text(f"UPDATE products SET {', '.join(sets)}, updated_at = NOW() WHERE id = :i"), p)
            await propagate_product_prices(p["i"], db)
            n += 1
    return n


@handler("PRICE_CHANGE")
class PriceChange:
    @staticmethod
    async def check(db, p: dict):
        if not p.get("changes"):
            raise HTTPException(400, "No price changes in this request.")

    @staticmethod
    async def post(db, req: dict, approver) -> str:
        n = await apply_prices(db, req["payload"]["changes"], approver.user_id, "PRICE_IMPORT", req["request_no"], req["reason"])
        if req["payload"].get("import_id"):
            await db.execute(text("UPDATE import_jobs SET status = 'IMPORTED', committed_at = NOW(), message = :m WHERE id = :i"),
                             {"m": f"{n} item price(s) updated after approval {req['request_no']}", "i": int(req["payload"]["import_id"])})
        return f"{n} items"
